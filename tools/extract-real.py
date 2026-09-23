"""
extract-real.py, menyusun src/data/org.ts dari dua berkas Excel yang dilampirkan.

Sumber:
  MasterDataArea (12).xlsx   sheet "Kanwil"  : master cabang, Kanwil induk, provinsi, kota
  Produksi up to Juli 2026.xlsx sheet "Sheet1": NWP dan profit 2025 dan 2026, kelas, target

Jalankan ulang bila kedua berkas diperbarui:
    python tools/extract-real.py "<path master>" "<path produksi>"

Catatan kepatuhan:
  Satu baris pada berkas produksi memuat nama perorangan pada kolom CABANG.
  Nama tersebut disamarkan menjadi kode unit, sedangkan angkanya dipertahankan
  agar total Kanwil tetap sesuai berkas sumber.
"""
import io
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\user\Downloads\MasterDataArea (12).xlsx'
PRODUKSI = sys.argv[2] if len(sys.argv) > 2 else r'C:\Users\user\Downloads\Produksi up to Juli 2026.xlsx'
OUT = os.path.join(ROOT, 'src', 'data', 'org.ts')

# Selisih penamaan antar kedua berkas. Kunci memakai nama pada berkas produksi.
ALIAS = {
    'BSD': 'BUMI SERPONG DAMAI',
    'PEKANBARU': 'PEKAN BARU',
    'SEMARANG AGENCY': 'AGENCY SEMARANG',
    'PEKAN BARU AGENCY': 'AGENCY PEKAN BARU',
    'PADANG SIDEMPUAN': 'PADANGSIDIMPUAN',
    'MANADO AGENCY': 'AGENCY MANADO',
    'MAKASSAR AGENCY': 'AGENCY MAKASSAR',
    'BATAM AGENCY': 'AGENCY BATAM',
}

# Nama perorangan yang muncul sebagai entri cabang, disamarkan.
SAMAR = {'CYNTHIA AGUSTINA': 'UNIT KHUSUS KANWIL 2'}


def norm(s):
    s = re.sub(r'[^A-Z0-9 ]', ' ', str(s).upper())
    return re.sub(r'\s+', ' ', s).strip()


def rapikan(nama):
    """Judul kota sederhana: buang prefiks KOTA / KAB."""
    if not nama:
        return ''
    n = str(nama).upper()
    n = re.sub(r'^(KOTA ADM\.?|KOTA|KAB\.?|KABUPATEN)\s+', '', n)
    return n.strip()


def main():
    wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    master = {}
    for r in wb['Kanwil'].iter_rows(min_row=2, values_only=True):
        if r[3]:
            master[norm(r[3])] = {'kanwil': str(r[1]).replace(' ', ''),
                                  'prov': str(r[6] or ''), 'kota': rapikan(r[7])}
    wb.close()

    wb = openpyxl.load_workbook(PRODUKSI, read_only=True, data_only=True)
    prod = [r for r in wb['Sheet1'].iter_rows(min_row=2, values_only=True) if r[2]]
    wb.close()

    baris, tanpa_geo = [], []
    for r in prod:
        nama_asli = str(r[2]).strip()
        nama = SAMAR.get(nama_asli.upper(), nama_asli).upper()
        kunci = norm(ALIAS.get(nama_asli.upper(), nama_asli))
        m = master.get(kunci)
        if not m:
            tanpa_geo.append(nama)
        kanwil = str(r[0]).replace(' ', '')
        no = int(r[1])
        kelas = str(r[3] or '').strip()
        nwp25 = float(r[4] or 0)
        target = float(r[6] or 0)
        nwp26 = float(r[7] or 0)
        profit25 = float(r[8] or 0)
        profit26 = float(r[13] or 0)
        baris.append({
            'kode': 'C%03d' % no, 'nama': nama, 'kanwil': kanwil, 'kelas': kelas,
            'prov': (m or {}).get('prov', ''), 'kota': (m or {}).get('kota', ''),
            'nwp25': nwp25, 'nwp26': nwp26, 'target': target,
            'profit25': profit25, 'profit26': profit26,
        })

    baris.sort(key=lambda b: (b['kanwil'], -b['nwp26']))

    def js(v):
        if isinstance(v, str):
            return "'" + v.replace("'", "\\'") + "'"
        if isinstance(v, float):
            return repr(int(round(v))) if abs(v - round(v)) < 0.5 else repr(round(v, 2))
        return repr(v)

    rows = ',\n'.join(
        '  [%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s]' % tuple(
            js(b[k]) for k in ('kode', 'nama', 'kanwil', 'kelas', 'prov', 'kota',
                               'nwp25', 'nwp26', 'target', 'profit25', 'profit26'))
        for b in baris)

    tot = {k: sum(b[k] for b in baris) for k in ('nwp25', 'nwp26', 'target', 'profit25', 'profit26')}

    isi = TEMPLATE.replace('__ROWS__', rows) \
                  .replace('__N__', str(len(baris))) \
                  .replace('__SUMBER_MASTER__', os.path.basename(MASTER)) \
                  .replace('__SUMBER_PRODUKSI__', os.path.basename(PRODUKSI))
    io.open(OUT, 'w', encoding='utf-8').write(isi)

    print('tertulis: %s' % OUT)
    print('cabang   : %d' % len(baris))
    for k, v in tot.items():
        print('%-9s: %.3f T' % (k, v / 1e12))
    if tanpa_geo:
        print('tanpa data geografis (%d): %s' % (len(tanpa_geo), ', '.join(tanpa_geo)))


TEMPLATE = u'''/* \u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550
   org.js, master cabang dan angka produksi.

   BERKAS INI DIHASILKAN OTOMATIS. Jangan disunting langsung.
   Sumber: __SUMBER_MASTER__ dan __SUMBER_PRODUKSI__
   Pembuat: tools/extract-real.py

   Isi yang bersumber dari data nyata Asuransi Sinar Mas:
     nama cabang, Kantor Wilayah induk, kelas cabang, provinsi, kota,
     NWP 2025, NWP 2026 sampai Juli, target NWP 2026, profit 2025 dan 2026.

   Satu entri memuat nama perorangan pada berkas sumber dan telah disamarkan
   menjadi UNIT KHUSUS KANWIL 2, sedangkan angkanya dipertahankan agar total
   Kantor Wilayah tetap sesuai berkas sumber.
   \u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550\u2550 */

const UNIT_KIND = { KANWIL: 'kanwil', DIVISI: 'divisi' }

/* Kantor Wilayah. Berkas yang tersedia baru mencakup jaringan cabang.
   Divisi seperti MBU, Health Insurance, Commercial Lines, Financial Insurance,
   Agency, dan ASNET belum tercakup karena datanya belum dilampirkan. Struktur
   di bawah siap menampung unit tersebut begitu datanya tersedia. */
const UNITS = [
  { id: 'KANWIL1', name: 'KANTOR WILAYAH 1', short: 'Kanwil 1', kind: UNIT_KIND.KANWIL },
  { id: 'KANWIL2', name: 'KANTOR WILAYAH 2', short: 'Kanwil 2', kind: UNIT_KIND.KANWIL },
  { id: 'KANWIL3', name: 'KANTOR WILAYAH 3', short: 'Kanwil 3', kind: UNIT_KIND.KANWIL },
]

/* Format: [kode, nama, kanwil, kelas, provinsi, kota,
            NWP 2025, NWP 2026 s.d. Juli, target NWP 2026, profit 2025, profit 2026] */
const BRANCHES_RAW: any[] = [
__ROWS__
]

const BRANCHES = BRANCHES_RAW.map(([code, name, unit, kelas, prov, kota,
  nwp25, nwp26, target, profit25, profit26]) => ({
  code, name, unit, kelas, prov, kota,
  nwp25, nwp26, target, profit25, profit26,
  id: 'B' + code,
  /* Beban mencakup klaim dan biaya, diturunkan dari selisih premi dan profit. */
  beban26: nwp26 - profit26,
  beban25: nwp25 - profit25,
  rasioBeban: nwp26 > 0 ? (nwp26 - profit26) / nwp26 : 0,
}))

/* \u2500\u2500 Enum aktivitas pemasaran \u2500\u2500
   Dibaca dari tampilan Matrix Distribution pada Pega ASMPro. Belum ada
   berkas ekspor untuk bagian ini, sehingga seluruh angka effort dan
   prospek pada prototipe ini merupakan simulasi. */
/* Tipe Matrix, mengikut Matrix Distribution. Porsinya dikalibrasi ke
   Tabel Summary Matrix Kanwil 1 yang dibaca 10 September 2026, yaitu
   2.193, 66, 38, 6, dan 48 dari 2.351 prospek. */
const SUMBER = ['New Prospek Marketing', 'Existing Prospek', 'Lapse Prospek', 'Tender', 'Crawling AI']
const SUMBER_MIX = [0.9328, 0.0281, 0.0162, 0.0026, 0.0204]

/* Status prospek, mengikut Matrix Distribution. Tiga yang pertama adalah
   pipeline yang masih hidup, tiga sisanya sudah tutup buku. Keterangan
   Cold, Warm, dan Hot dikutip apa adanya dari layar Pega.

   Porsi dikalibrasi ke Kanwil 1: 208, 118, 58, 1.422, 25, dan 520.
   Perhatikan bahwa pipeline hidup hanya 16,3 persen dari seluruh prospek. */
const STATUS = ['Cold', 'Warm', 'Hot', 'Terbit Polis', 'Pending', 'Batal']
const STATUS_MIX = [0.0885, 0.0502, 0.0247, 0.6048, 0.0106, 0.2212]
const STATUS_HIDUP = ['Cold', 'Warm', 'Hot']
const STATUS_KET = {
  Cold: 'Perkenalan dan komunikasi awal, atau kirim penawaran',
  Warm: 'Sudah komunikasi lebih lanjut (follow up) atau presentasi dan penjelasan produk',
  Hot: 'Sudah mengirimkan proposal penawaran atau TC',
  'Terbit Polis': 'Prospek sudah menjadi polis',
  Pending: 'Tertahan, menunggu pihak lain',
  Batal: 'Tidak jadi',
}

/* Status effort terakhir, dimensi yang berbeda dari status prospek di atas.
   Nilainya terbaca pada tabel Belum Melakukan Follow Up di Pega. Dimensi ini
   menerangkan langkah terakhir yang dikerjakan, bukan panas dinginnya prospek. */
const TAHAP = ['Prospek Baru', 'Tahap Awal', 'Follow Up', 'Negosiasi', 'Proposal']

const LINI = ['MBU', 'Personal Accident', 'Aneka', 'Marine Cargo', 'Travel', 'Fire', 'HID']
const LINI_MIX = [0.2, 0.14, 0.13, 0.15, 0.07, 0.21, 0.1]

/* Kanal distribusi.

   Materi training ASM (MDP Impact, Transkrip Ringkas d16) menyebut tujuh
   kategori Sumber Bisnis dengan pembelahan resmi Direct dan Indirect.
   Pengelompokan di bawah mengikuti keputusan pengguna dan berbeda pada dua
   titik, yang keduanya dicatat di README:

   1. Kelompok kedua dinamai Captive. Pada Notulen Rapim, kata captive dipakai
      untuk asuransi milik grup mitra yang justru mengambil bisnis ASM, misalnya
      captive BCAI dan leasing captive ACC. Keterangan pembeda ditampilkan di layar.
   2. Agency ditempatkan pada kelompok Direct, sedangkan materi training
      menempatkannya pada Indirect. */
const KANAL = {
  DIRECT: { id: 'DIRECT', label: 'Direct', sub: ['Direct', 'Corporate', 'Agency'] },
  CAPTIVE: { id: 'CAPTIVE', label: 'Captive', sub: ['Banking', 'Multifinance', 'Broker', 'Sinar Mas Group'] },
}
const SUB_KANAL = [...KANAL.DIRECT.sub, ...KANAL.CAPTIVE.sub]
const KANAL_DARI_SUB = SUB_KANAL.reduce((acc, s) => {
  acc[s] = KANAL.DIRECT.sub.includes(s) ? 'DIRECT' : 'CAPTIVE'
  return acc
}, {})

const MAPPING_KAT = ['Baru', 'Extension', 'New Business', 'Renewal Business']

/* Bobot indeks effort masih bersifat usulan dan perlu disepakati sebelum
   digunakan sebagai dasar penilaian kinerja. */
const AKTIVITAS = [
  { id: 'kunjungan', label: 'Kunjungan', w: 3 },
  { id: 'penawaran', label: 'Penawaran', w: 2.5 },
  { id: 'follow_up', label: 'Follow Up', w: 1.5 },
  { id: 'telp', label: 'Telepon', w: 1 },
]

/* Nama Marketing Officer dan nama badan usaha di bawah ini seluruhnya
   rekaan. Tidak ada data karyawan maupun nasabah yang digunakan. */
const NAMA_DEPAN = [
  'Andi', 'Budi', 'Citra', 'Dian', 'Eka', 'Fajar', 'Gita', 'Hendra', 'Indra', 'Joko',
  'Kartika', 'Lestari', 'Mira', 'Nanda', 'Oki', 'Putri', 'Rahmat', 'Sari', 'Tono', 'Umi',
  'Vina', 'Wahyu', 'Yuda', 'Zahra', 'Bagus', 'Deni', 'Erni', 'Fitri', 'Galih', 'Hesti',
  'Ilham', 'Jihan', 'Kurnia', 'Lukman', 'Maya', 'Nurul', 'Okta', 'Prima', 'Rico', 'Silvi',
]
const NAMA_BELAKANG = [
  'Prasetyo', 'Wijaya', 'Santoso', 'Hidayat', 'Nugroho', 'Kusuma', 'Halim', 'Saputra',
  'Ramadhan', 'Permana', 'Setiawan', 'Anggraini', 'Maulana', 'Firmansyah', 'Yulianti',
  'Wibowo', 'Susanto', 'Handoko', 'Rahayu', 'Utami', 'Gunawan', 'Marpaung', 'Sitorus',
  'Simanjuntak', 'Lubis', 'Tanjung', 'Hasibuan', 'Pratama', 'Mahendra', 'Iskandar',
]
const PT_DEPAN = [
  'Karya', 'Mitra', 'Sentosa', 'Bahari', 'Cakra', 'Delta', 'Eka', 'Graha', 'Harmoni',
  'Indo', 'Jaya', 'Kencana', 'Lintas', 'Mandala', 'Nusa', 'Optima', 'Perkasa', 'Rajawali',
  'Surya', 'Tirta', 'Utama', 'Wana', 'Bumi', 'Cipta', 'Dwi', 'Anugerah', 'Berkah',
]
const PT_BELAKANG = [
  'Sejahtera', 'Abadi', 'Makmur', 'Nusantara', 'Persada', 'Mandiri', 'Lestari', 'Utama',
  'Prima', 'Jaya', 'Sentosa', 'Gemilang', 'Bersama', 'Selaras', 'Perdana', 'Mulia',
]
const PT_BIDANG = [
  'Logistik', 'Manufaktur', 'Konstruksi', 'Energi', 'Agro', 'Properti', 'Otomotif',
  'Farmasi', 'Tekstil', 'Pangan', 'Kimia', 'Teknologi', 'Pelayaran', 'Ritel',
]

/* Hari libur nasional 2026 menurut perkiraan, dipakai agar deret harian
   memiliki lembah yang dapat dijelaskan secara kalender. */
const LIBUR_2026 = new Set([
  '2026-01-01', '2026-01-17', '2026-02-17', '2026-03-19', '2026-03-20', '2026-03-21',
  '2026-03-22', '2026-03-23', '2026-04-03', '2026-05-01', '2026-05-14', '2026-05-26',
  '2026-05-27', '2026-05-28', '2026-05-29', '2026-06-01',
])

const ORG = {
  UNITS, BRANCHES, UNIT_KIND,
  SUMBER, SUMBER_MIX, TAHAP, STATUS, STATUS_MIX, STATUS_HIDUP, STATUS_KET, LINI, LINI_MIX, MAPPING_KAT, AKTIVITAS,
  KANAL, SUB_KANAL, KANAL_DARI_SUB,
  NAMA_DEPAN, NAMA_BELAKANG, PT_DEPAN, PT_BELAKANG, PT_BIDANG, LIBUR_2026,
  JUMLAH_CABANG: __N__,

  unitById: (id) => UNITS.find((u) => u.id === id),
  branchById: (id) => BRANCHES.find((b) => b.id === id),
  branchesOf: (unitId) => BRANCHES.filter((b) => b.unit === unitId),
}
'''

if __name__ == '__main__':
    main()
